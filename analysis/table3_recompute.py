"""M2/M3/N3/N4: recompute the variable-equivalence component of Table 3 on the CURRENT noWorkflow 2.1
graphs, with one rule for every version (the rule of Figure 8): a ground-truth data item counts as a
true positive when a node labeled with its mapped name (column of mappings.xlsx for that version)
exists in the .dot; one graph node can satisfy at most one ground-truth item (greedy, in sheet
order); the seven YesWorkflow ports are always false negatives; FP_var = universe - TP_var, where
the universe is, by default, every node of the graph except function-call nodes and
definition-reference nodes (--universe data), or every node except definition-reference nodes
(--universe all, the convention of Figure 8 and of the original spreadsheet, where library-call
nodes of the monolithic graph count as false positives).

Usage: python table3_recompute.py <mappings.xlsx> <version>=<graph.dot>[:<script.py>] [...]
       [--universe data|all] [--ref <expert.py>] [--aliases <aliases.csv>] [--export <prefix>]
       [--db <version>=<.noworkflow/db.sqlite> ...]
Option values are never read as version specs; unknown versions or options abort. With --db, every
alias requires --db for its version AND for expert, and is verified against the noWorkflow databases:
the node id e_<N> is the evaluation id; its repr (final value recorded by noWorkflow) must equal the
manifest's value_version, and the repr of the expert node automatically matched to the same
ground-truth item must equal value_gt and value_version. Aliases on call or definition nodes are
rejected. Without --db the manifest is applied as declared and the output says "UNVERIFIED".
Two variants are always reported when --aliases is given: (A) automatic, by label only, and (B)
automatic + value-based aliases from the manifest. The manifest (one row per manual match) has the
columns: version, gt_item (YW name, column A of sheet Nodes), node_id (e_... in the .dot), node_label,
value_gt, value_version, evidence (how the equality of final values was established, e.g. the
noWorkflow evaluation ids). Invariants enforced: a node is used by at most one item, an item by at
most one node, the node must exist in the .dot and must not already be matched automatically; a
violated invariant aborts. --export writes <prefix>_<version>.csv with every ground-truth item, its
status (TP-auto, TP-alias, FN) and the node id, so that the trail node <-> item is auditable.
  (the script is used to recognize definition-reference nodes, one per user-defined function)
  version in {monolithic, expert, gemini}; e.g.
  python table3_recompute.py mappings.xlsx monolithic=monolithic.dot expert=modular_expert.dot gemini=modular_llm.dot
Functional component: computed live with compare_functions_lib (metric Sim of Section 4.6.2) when
--ref <expert script> is given and the version's script is given after ':' (the expert itself gives
14/14/14; the monolithic has no functions, 0/0/0). Without --ref, the values of similarity.txt are
used for gemini (AI->G 6.84, G->AI 9.44, |AI| 8) and printed as such. |G| = 14.
Matching policy: label_match.resolve_items (one graph node satisfies at most one ground-truth
item, in sheet order; when several nodes carry the item's name, the stage rule keeps the node that
depends on a call of the item's producer block in sheet Edges; an item left with several candidates
counts as TP but is exported as TP-ambiguous with every candidate, never as an asserted identity).
In the expert graph `delp` and `delt` have two nodes each (190: read from the file; 206: adjusted),
assigned by the stage rule to InputDelP/InputDelT (producer ReadDataFileInput) and DelP/DelT
(producer CalculateRefsGeometryModel).
Prints, per version: list of matched / unmatched items, TP_var, FN_var, FP_var, and SP, SR, SF.
Label matching follows label_match.py (exact names, one delimited path rule, no suffixes, no
abbreviated paths, no case folding; audit G1-F02, G1-F03); the nodes
eligible for a match are exactly the nodes of the variable universe (never call or
definition-reference nodes), and 0 <= TP_var <= universe, FP_var >= 0 are enforced (audit G1-F01,
G1-F02). With --ref, every version must give its script after ':' (otherwise the run aborts);
without --ref, each version's functional component is tagged as the similarity.txt fallback, which
is not certified (audit G1-R01).
"""
import ast, os, re, sys
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from label_match import CALL, body, function_of, producers, resolve_items

COL = {"monolithic": 2, "expert": 4, "gemini": 8}   # 0-based columns of the mapped name in sheet Nodes
FUNC = {"monolithic": (0.0, 0.0, 0), "expert": (14.0, 14.0, 14), "gemini": (6.84, 9.44, 8)}
G = 14

def load_items(xlsx):
    ws = openpyxl.load_workbook(xlsx)["Nodes"]
    items = []
    for r in ws.iter_rows(min_row=21, max_row=60, values_only=True):
        if r[0] and r[10] in ("VAR", "INPUT", "OUTPUT"):
            items.append(r)
    return items

def parse_dot(path):
    nodes = []
    for line in open(path, errors="replace"):
        m = re.match(r'\s*([ea]_\w+)\s*\[label="(.*?)"\s+fillcolor="([^"]*)"', line)
        if m:
            nodes.append((m.group(1), m.group(2), m.group(3)))
    return nodes

def parse_edges(path):
    return {(s, d) for line in open(path, errors="replace") for s, d in re.findall(r'([ea]_\w+)\s*->\s*([ea]_\w+)', line)}

def main():
    xlsx = sys.argv[1]
    # positional <version>=<graph.dot>[:<script.py>] specs; the value after an option is never a spec
    specs, dbs, i = [], {}, 2
    while i < len(sys.argv):
        a = sys.argv[i]
        if a in ("--universe", "--ref", "--aliases", "--export", "--db"):
            assert i + 1 < len(sys.argv), f"{a} needs a value"
            if a == "--db":
                assert "=" in sys.argv[i + 1], f"--db needs <version>=<db.sqlite>, got {sys.argv[i + 1]}"
                v, path = sys.argv[i + 1].split("=", 1); dbs[v] = path
            i += 2; continue
        assert not a.startswith("--"), f"unknown option {a}"
        assert "=" in a, f"not a <version>=<graph.dot>[:<script.py>] spec: {a}"
        specs.append(a); i += 1
    for v in [s.split("=", 1)[0] for s in specs] + list(dbs):
        assert v in COL, f"unknown version {v}: expected one of {sorted(COL)}"
    universe_mode = sys.argv[sys.argv.index("--universe") + 1] if "--universe" in sys.argv else "data"
    ref_script = sys.argv[sys.argv.index("--ref") + 1] if "--ref" in sys.argv else None
    aliases_path = sys.argv[sys.argv.index("--aliases") + 1] if "--aliases" in sys.argv else None
    export = sys.argv[sys.argv.index("--export") + 1] if "--export" in sys.argv else None
    def db_repr(version, node_id):
        import sqlite3
        c = sqlite3.connect(dbs[version])
        r = c.execute("select repr from evaluation where id=?", (int(node_id.split("_")[1]),)).fetchone()
        return None if r is None else r[0]
    expert_trail = {}
    aliases = []
    if aliases_path:
        import csv
        aliases = list(csv.DictReader(open(aliases_path, encoding="utf-8")))
    if ref_script:
        missing = [s.split("=", 1)[0] for s in specs if ":" not in s.split("=", 1)[1]]
        assert not missing, f"--ref given but no script for {missing}: give <version>=<graph.dot>:<script.py> for every version"
        from compare_functions_lib import extract_functions, function_similarity
        ref_funcs = extract_functions(ref_script)
        assert len(ref_funcs) == G, f"reference script has {len(ref_funcs)} functions, expected {G}"
    items = load_items(xlsx)
    assert len(items) == 40, len(items)
    for spec in specs:
        version, rest = spec.split("=", 1)
        dot, script = (rest.split(":", 1) + [None])[:2]
        nodes = parse_dot(dot)
        calls = [n for n in nodes if n[2] == CALL]
        if not script:
            print(f"WARNING {version}: no script given, definition-reference nodes cannot be recognized")
        funcs = {f.name for f in ast.walk(ast.parse(open(script, encoding="utf-8").read())) if isinstance(f, ast.FunctionDef)} if script else set()
        defrefs = [n for n in nodes if n[2] != CALL and not n[0].startswith("a_") and function_of(n[1], funcs)]
        excluded = {x[0] for x in calls} | {x[0] for x in defrefs}
        eligible = [n for n in nodes if n[0] not in excluded]   # = the variable universe of --universe data
        res = resolve_items([(it[0], it[COL[version]]) for it in items], [(n, l) for n, l, c in eligible],
                            parse_edges(dot), producers(openpyxl.load_workbook(xlsx), nodes, funcs))
        matched, unmatched, trail = [], [], {}
        for it in items:
            st, cands = res[it[0]]
            if cands:
                matched.append((it[0], it[COL[version]])); trail[it[0]] = (st, cands[0])
            else:
                unmatched.append((it[0], it[COL[version]]))
        used = {n for st, cands in res.values() for n in cands}   # an alias may not take a candidate of a matched item
        if version == "expert":
            expert_trail = dict(trail)
        # variant B: value-based aliases from the manifest (checked, never guessed)
        alias_matched = []
        if not [a for a in aliases if a["version"] == version]:
            alias_status = "NONE: no alias for this version in the manifest, nothing to verify"
        else:
            alias_status = "UNVERIFIED (no --db)"
        ids = {nid for nid, l, c in nodes}
        callids = excluded
        for a in [a for a in aliases if a["version"] == version]:
            assert a["gt_item"] in {it[0] for it in items}, f"alias: unknown ground-truth item {a['gt_item']}"
            assert a["node_id"] in ids, f"alias: node {a['node_id']} not in {dot}"
            assert a["node_id"] not in callids, f"alias: node {a['node_id']} is a call or definition node"
            assert a["node_id"] not in used, f"alias: node {a['node_id']} already matched"
            assert a["gt_item"] not in trail, f"alias: item {a['gt_item']} already matched"
            assert a["value_gt"] == a["value_version"] and a["evidence"], f"alias {a['gt_item']}: values differ or no evidence"
            assert not dbs or (version in dbs and "expert" in dbs), f"alias {a['gt_item']}: --db given, so --db {version}=... and --db expert=... are both required"
            if dbs:
                rv = db_repr(version, a["node_id"])
                assert rv is not None, f"alias {a['gt_item']}: node {a['node_id']} has no evaluation in {dbs[version]}"
                assert rv == a["value_version"], f"alias {a['gt_item']}: DB value {rv!r} != manifest {a['value_version']!r}"
                gt_node = expert_trail.get(a["gt_item"], ("", ""))[1]
                assert gt_node, f"alias {a['gt_item']}: expert must be processed first (give expert=... before {version}=...) and the item must be matched there"
                rg = db_repr("expert", gt_node)
                assert rg == a["value_gt"] == rv, f"alias {a['gt_item']}: expert DB value {rg!r} != manifest/version value {a['value_gt']!r}/{rv!r}"
                alias_status = "VERIFIED against noWorkflow databases"
            used.add(a["node_id"]); alias_matched.append((a["gt_item"], a["node_id"])); trail[a["gt_item"]] = ("TP-alias", a["node_id"])
            unmatched = [u for u in unmatched if u[0] != a["gt_item"]]
        for it in items:
            trail.setdefault(it[0], ("FN", ""))
        if export:
            import csv
            with open(f"{export}_{version}.csv", "w", newline="") as f:
                w = csv.writer(f); w.writerow(["gt_item", "mapped_name", "status", "node_id", "node_label", "candidates"])
                lab = {nid: l for nid, l, c in nodes}
                for it in items:
                    st, nid = trail[it[0]]
                    w.writerow([it[0], it[COL[version]], st, nid, lab.get(nid, ""), "|".join(res[it[0]][1]) if st == "TP-ambiguous" else ""])
        tp_auto = len(matched)
        for variant, tp in (("A automatic", tp_auto), ("B automatic + aliases", tp_auto + len(alias_matched))):
          if variant.startswith("B") and not aliases_path: continue   # an empty manifest still reports B (= A)
          fn = (40 - tp) + 7
          universe = len(eligible) + (len(calls) if universe_mode == "all" else 0)
          fp = universe - tp
          assert 0 <= tp <= len(eligible) <= universe and fp >= 0, f"{version}: TP_var={tp} outside the universe {universe}"
          tp_ai_g, tp_g_ai, n_ai = FUNC[version]
          func_src = "FALLBACK: similarity.txt literal values, NOT CERTIFIED"
          if ref_script:
              cand = extract_functions(script)
              n_ai = len(cand)
              tp_g_ai = sum(max((function_similarity(lg, lc) for lc in cand.values()), default=0.0) for lg in ref_funcs.values())
              tp_ai_g = sum(max((function_similarity(lc, lg) for lg in ref_funcs.values()), default=0.0) for lc in cand.values())
              func_src = f"computed from {ref_script} x {script}"
          fp_f = n_ai - tp_ai_g; fn_f = G - tp_g_ai
          sp = (tp_ai_g + tp) / ((tp_ai_g + fp_f) + (tp + fp)) if (tp + fp + n_ai) else 0
          sr = (tp_g_ai + tp) / ((tp_g_ai + fn_f) + (tp + fn))
          sf = 2 / (1 / sp + 1 / sr) if sp and sr else 0
          print(f"\n=== {version} [{variant}]: {dot}: {len(nodes)} nodes, {len(calls)} call nodes, {len(defrefs)} definition-reference nodes, universe ({universe_mode}) {universe}")
          print(f"TP_var={tp}  FN_var={fn} (incl. 7 ports)  FP_var={fp}   TP_func AI->G={tp_ai_g:.2f} G->AI={tp_g_ai:.2f} |AI|={n_ai} FP_func={fp_f:.2f} FN_func={fn_f:.2f}  [{func_src}]")
          print(f"SP={sp:.3f}  SR={sr:.3f}  SF={sf:.3f}")
          print("matched (auto) :", ", ".join(f"{a}->{b} [{trail[a][1]}{' AMBIGUOUS: ' + '|'.join(res[a][1]) if trail[a][0] == 'TP-ambiguous' else ''}]" for a, b in matched))
          if variant.startswith("B"): print(f"matched (alias) [{alias_status}]:", ", ".join(f"{a}->{b}" for a, b in alias_matched))
          print("unmatched      :", ", ".join(f"{a}->{b}" for a, b in unmatched))

if __name__ == "__main__":
    main()
