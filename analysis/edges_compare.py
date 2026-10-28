"""M4: edge-level comparison between the YesWorkflow ground truth and a noWorkflow graph, with an
explicit unit of comparison and a one-to-one matching.

Ground truth = the canonical set of YesWorkflow edges: rows INPUT/OUTPUT of sheet Edges of
mappings.xlsx with "Exists in YW" = YES, deduplicated by (type, YW from, YW to); duplicated rows are
reported. Candidate graph edges = edges of the .dot whose two endpoints are in the compared node set
(function-call nodes of the blocks, file nodes, and data nodes that the sheet Nodes maps to
ground-truth items). A ground-truth edge (data item <-> block) is compatible with a graph edge when
the graph edge joins a call node of that block and a node of that data item, in either orientation
(noWorkflow draws dependent -> dependency, YesWorkflow draws data flow). A maximum one-to-one
matching between ground-truth edges and candidate edges gives TP; FN = ground-truth edges left
unmatched; FP = candidate edges left unmatched. Precision = TP/(TP+FP), recall = TP/(TP+FN).
Edges touching nodes outside the compared set are listed as "not comparable" and not counted.
Label matching follows label_match.py, the same rule as table3_recompute.py (exact names, one
delimited path rule, no suffixes; audit G1-F02, G1-F03); definition-reference nodes are never data
nodes (G1-F01). The data endpoint of a ground-truth edge is the node that label_match.resolve_items
assigns to the item in sheet Nodes, the same trail as table3_recompute (G1-F04); an item that keeps
several candidates is exported as TP-ambiguous, never as an asserted identity; an item absent from
sheet Nodes (PlotFile) is resolved by the sheet Edges name and flagged. The matching is a
deterministic augmenting-path maximum matching (ground-truth edges in sheet order; for each, the
graph edges in noWorkflow's orientation first, then the reversed ones), independent of
PYTHONHASHSEED; the counts do not depend on this order, the exported identity does.

Usage: python edges_compare.py <mappings.xlsx> <graph.dot> <script.py> [--version monolithic|expert|gemini] [--export out.csv]
"""
import ast, csv, os, re, sys
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from label_match import CALL, block_calls, body, data_match, function_of, producers, resolve_items

COL = {"monolithic": 2, "expert": 4, "gemini": 8}

def main():
    xlsx, dot, script = sys.argv[1:4]
    version = sys.argv[sys.argv.index("--version") + 1] if "--version" in sys.argv else ("expert" if "expert" in dot else "gemini" if ("llm" in dot or "gemini" in dot) else "monolithic")
    export = sys.argv[sys.argv.index("--export") + 1] if "--export" in sys.argv else None
    funcs = {f.name for f in ast.walk(ast.parse(open(script, encoding="utf-8").read())) if isinstance(f, ast.FunctionDef)}
    nodes, edges = {}, set()
    for line in open(dot, errors="replace"):
        m = re.match(r'\s*([ea]_\w+)\s*\[label="(.*?)"\s+fillcolor="([^"]*)"', line)
        if m:
            nodes[m.group(1)] = (m.group(2), m.group(3)); continue
        for s, d in re.findall(r'([ea]_\w+)\s*->\s*([ea]_\w+)', line):
            edges.add((s, d))
    defrefs = {n for n, (l, c) in nodes.items() if c != CALL and not n.startswith("a_") and function_of(l, funcs)}
    nlist = [(n, l, c) for n, (l, c) in nodes.items()]
    eligible = [(n, l) for n, l, c in nlist if c != CALL and n not in defrefs]
    def data_nodes(name):
        return [n for n, l in eligible if data_match(name, n, l)]
    wb = openpyxl.load_workbook(xlsx)
    # data endpoint of every ground-truth item: the same item -> node trail as table3_recompute
    items = [(r[0], r[COL[version]]) for r in wb["Nodes"].iter_rows(min_row=21, max_row=60, values_only=True) if r[0] and r[10] in ("VAR", "INPUT", "OUTPUT")]
    res = resolve_items(items, eligible, edges, producers(wb, nlist, funcs))
    # canonical ground-truth edge set
    gt, dups, outside = {}, [], set()
    for r in wb["Edges"].iter_rows(min_row=5, values_only=True):
        if r[0] in ("INPUT", "OUTPUT") and r[3] == "YES":
            key = (r[0], r[1], r[2])
            if key in gt: dups.append(key); continue
            item, block = (r[1], r[2]) if r[0] == "INPUT" else (r[2], r[1])
            fn = (r[5] if r[0] == "INPUT" else r[4]) or block
            if item in res:
                st, ds = res[item]
            else:   # item absent from sheet Nodes: data endpoint by the sheet Edges name, flagged
                outside.add(item)
                ds = data_nodes((r[4] if r[0] == "INPUT" else r[5]) or item)
                st = "TP-auto" if len(ds) == 1 else "TP-ambiguous" if ds else "FN"
            gt[key] = (fn, ds, st)
    # compared node set: call nodes, file nodes, data endpoints of the ground-truth items
    compared = {n for n, l, c in nlist if c == CALL or n.startswith("a_") or c == "white"}
    for fn, ds, st in gt.values():
        compared.update(ds)
    cand = sorted(e for e in edges if e[0] in compared and e[1] in compared)
    cset = set(cand)
    # compatibility (orientation ignored, by construction) and deterministic maximum matching
    adj = {}
    for key, (fn, ds, st) in gt.items():
        # noWorkflow orientation first (INPUT: call -> data; OUTPUT: data -> call), the reverse only after it
        fwd = [(a, b) if key[0] == "INPUT" else (b, a) for a in block_calls(nlist, funcs, fn) for b in ds]
        adj[key] = [e for e in fwd if e in cset] + [(e[1], e[0]) for e in fwd if (e[1], e[0]) in cset]
    owner = {}
    def augment(key, seen):
        for e in adj[key]:
            if e not in seen:
                seen.add(e)
                if e not in owner or augment(owner[e], seen):
                    owner[e] = key; return True
        return False
    for key in gt:
        augment(key, set())
    matched = {k: e for e, k in owner.items()}
    tp = len(matched); fn_list = [k for k in gt if k not in matched]
    fp_edges = [e for e in cand if e not in owner]
    p = tp / (tp + len(fp_edges)) if tp + len(fp_edges) else 0.0; rc = tp / len(gt) if gt else 0.0
    print(f"{dot} [{version}]: ground-truth edges={len(gt)} (duplicated sheet rows ignored: {len(dups)}); graph edges={len(edges)}, comparable={len(cand)}, not comparable={len(edges) - len(cand)}")
    print(f"one-to-one matching: TP={tp}  FN={len(fn_list)}  FP={len(fp_edges)}  precision={p:.3f}  recall={rc:.3f}  F={2*p*rc/(p+rc) if p+rc else 0:.3f}")
    for d in dups: print("  duplicate sheet row:", d)
    for i in sorted(outside): print("  item not in sheet Nodes, data endpoint by the sheet Edges name:", i, [x for x in gt.items() if i in x[0]][0][1][1])
    amb = [k for k in matched if gt[k][2] == "TP-ambiguous"]
    for k in amb: print("  TP-ambiguous (count only, identity not asserted):", k, "candidates", gt[k][1])
    for k in fn_list: print("  FN:", k, "->", gt[k][:2])
    for e in fp_edges: print("  FP:", e, body(nodes[e[0]][0]), "->", body(nodes[e[1]][0]))
    if export:
        with open(export, "w", newline="") as f:
            w = csv.writer(f); w.writerow(["status", "type", "yw_from", "yw_to", "graph_edge", "label_from", "label_to"])
            for k in gt:
                if k in matched:
                    e = matched[k]; w.writerow(["TP-ambiguous" if gt[k][2] == "TP-ambiguous" else "TP", *k, e, body(nodes[e[0]][0]), body(nodes[e[1]][0])])
            for k in fn_list: w.writerow(["FN", *k, "", "", ""])
            for e in fp_edges: w.writerow(["FP", "", "", "", e, body(nodes[e[0]][0]), body(nodes[e[1]][0])])

if __name__ == "__main__":
    main()
